"""Báo cáo AvT và các báo cáo quản lý (5.10)."""

import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import clock, domain as d
from app.models import (
    AuditLog,
    CountLine,
    CountSession,
    Handover,
    HandoverLine,
    Ingredient,
    Investigation,
    MenuItem,
    OrderItem,
    Receipt,
    ReceiptLine,
    Shift,
    StockTxn,
    User,
    WasteSlip,
)
from app.services import variance
from app.services.common import dec, money, q3


def _range(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    return clock.day_bounds(date_from)[0], clock.day_bounds(date_to)[1]


def _names(db: Session) -> dict[int, str]:
    return {u.id: u.full_name for u in db.scalars(select(User))}


def avt(
    db: Session, date_from: date, date_to: date, location: str | None = None, ingredient_id: int | None = None,
    shift_id: int | None = None, by_shift: bool = False,
) -> dict:
    """Lý thuyết, thực tế, hao hụt đã ghi, chênh lệch chưa giải thích (g, %, tiền) theo nguyên liệu,
    vị trí và (tùy chọn) ca; sau điều tra hiển thị phần đã giải trình theo nguyên nhân."""
    start, end = _range(date_from, date_to)
    ings = {i.id: i for i in db.scalars(select(Ingredient))}
    shifts_by_id = {s.id: s for s in db.scalars(select(Shift).where(Shift.planned_start < end + timedelta(days=1),
                                                                    Shift.planned_end > start - timedelta(days=1)))}
    names = _names(db)
    rows: dict[tuple, dict] = {}

    def key_of(ing_id, loc, sh):
        return (ing_id, loc, sh if by_shift else None)

    def row(ing_id, loc, sh):
        k = key_of(ing_id, loc, sh)
        if k not in rows:
            rows[k] = {"theoretical": Decimal(0), "waste": Decimal(0), "unexplained": Decimal(0),
                       "value": Decimal(0), "explained": defaultdict(lambda: [Decimal(0), Decimal(0)]),
                       "pending": Decimal(0), "counts": 0}
        return rows[k]

    tx = select(StockTxn.ingredient_id, StockTxn.location, StockTxn.shift_id, StockTxn.txn_type, func.sum(StockTxn.qty)) \
        .where(StockTxn.occurred_at >= start, StockTxn.occurred_at < end,
               StockTxn.txn_type.in_([*d.THEORETICAL_TYPES, *d.RECORDED_WASTE_TYPES]),
               StockTxn.late.is_(False)) \
        .group_by(StockTxn.ingredient_id, StockTxn.location, StockTxn.shift_id, StockTxn.txn_type)
    cl = select(CountLine).where(CountLine.cutoff_at >= start, CountLine.cutoff_at < end, CountLine.variance_qty.is_not(None))
    if location:
        tx = tx.where(StockTxn.location == location)
        cl = cl.where(CountLine.location == location)
    if ingredient_id:
        tx = tx.where(StockTxn.ingredient_id == ingredient_id)
        cl = cl.where(CountLine.ingredient_id == ingredient_id)
    if shift_id:
        tx = tx.where(StockTxn.shift_id == shift_id)
        cl = cl.where(CountLine.shift_id == shift_id)

    for ing_id, loc, sh, ttype, total in db.execute(tx):
        r = row(ing_id, loc, sh)
        if ttype in d.THEORETICAL_TYPES:
            r["theoretical"] -= dec(total)
        else:
            r["waste"] -= dec(total)

    lines = list(db.scalars(cl))
    invs = {i.count_line_id: i for i in db.scalars(
        select(Investigation).where(Investigation.count_line_id.in_([ln.id for ln in lines] or [0])))}
    for ln in lines:
        r = row(ln.ingredient_id, ln.location, ln.shift_id)
        r["unexplained"] += ln.variance_qty
        r["value"] += ln.variance_value or 0
        r["counts"] += 1
        inv = invs.get(ln.id)
        if inv:
            for a in inv.allocations:
                r["explained"][a["cause"]][0] += dec(a["qty"])
                r["explained"][a["cause"]][1] += dec(a["value"])
            if inv.status == "open":
                r["pending"] += inv.variance_qty - inv.settled_qty

    out = []
    for (ing_id, loc, sh), r in rows.items():
        ing = ings[ing_id]
        theo, unexpl = q3(r["theoretical"]), q3(r["unexplained"])
        tier, pct = variance.classify(unexpl, theo, variance.thresholds_for(db, ing))
        s = shifts_by_id.get(sh) or (db.get(Shift, sh) if sh else None)
        out.append({
            "ingredient_id": ing_id, "ingredient": ing.name, "group": ing.control_group, "unit": ing.base_unit,
            "location": loc, "shift_id": sh,
            "shift": f"{s.name} {s.business_date:%d/%m}" if s else None,
            "head_chef": names.get(s.head_chef_id) if s else None,
            "theoretical": theo, "recorded_waste": q3(r["waste"]),
            "actual": q3(theo + r["waste"] + unexpl), "unexplained": unexpl, "pct": pct,
            "value": money(r["value"]), "tier": tier if unexpl != 0 else d.TIER_OK, "counts": r["counts"],
            "explained": [
                {"cause": c, "label": d.CAUSES[c], "qty": q3(v[0]), "value": money(v[1])}
                for c, v in r["explained"].items()
            ],
            "pending_qty": q3(r["pending"]),
        })
    out.sort(key=lambda x: (-abs(x["value"]), x["ingredient"]))
    return {
        "from": date_from, "to": date_to, "rows": out,
        "total_value": money(sum((x["value"] for x in out), Decimal(0))),
        "total_shortage_value": money(sum((x["value"] for x in out if x["value"] > 0), Decimal(0))),
    }


def avt_excel(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "AvT"
    ws.append([f"Báo cáo thực tế so với lý thuyết (AvT) từ {report['from']:%d/%m/%Y} đến {report['to']:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    header = ["Nguyên liệu", "Nhóm", "Vị trí", "Ca", "Bếp trưởng ca", "Đơn vị", "Lý thuyết", "Hao hụt đã ghi",
              "Thực tế", "Chênh lệch chưa giải thích", "%", "Giá trị (đ)", "Mức", "Đã giải trình", "Chờ giải trình"]
    ws.append(header)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in report["rows"]:
        ws.append([
            r["ingredient"], r["group"], d.LOCATIONS[r["location"]], r["shift"] or "", r["head_chef"] or "",
            r["unit"], float(r["theoretical"]), float(r["recorded_waste"]), float(r["actual"]),
            float(r["unexplained"]), float(r["pct"]) if r["pct"] is not None else None, float(r["value"]),
            d.TIERS[r["tier"]],
            "; ".join(f"{e['label']}: {float(e['qty']):g} ({float(e['value']):,.0f} đ)" for e in r["explained"]),
            float(r["pending_qty"]) or None,
        ])
    ws.append([])
    ws.append(["Tổng giá trị chênh lệch", "", "", "", "", "", "", "", "", "", "", float(report["total_value"])])
    widths = [22, 7, 8, 16, 18, 8, 12, 14, 12, 16, 8, 14, 16, 40, 14]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=4, min_col=7, max_col=12):
        for c in row:
            c.number_format = "#,##0.##"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def waste_report(db: Session, date_from: date, date_to: date) -> dict:
    """Hao hụt theo loại, món, ca, nhân viên; tách ghi đúng lúc và ghi muộn qua điều tra."""
    start, end = _range(date_from, date_to)
    names = _names(db)
    txns = db.execute(
        select(StockTxn, Ingredient).join(Ingredient, Ingredient.id == StockTxn.ingredient_id)
        .where(StockTxn.txn_type.in_(d.RECORDED_WASTE_TYPES), StockTxn.occurred_at >= start, StockTxn.occurred_at < end)
    ).all()
    slips = {s.id: s for s in db.scalars(select(WasteSlip).where(WasteSlip.at >= start - timedelta(days=1), WasteSlip.at < end))}
    dish_of = dict(db.execute(select(OrderItem.id, MenuItem.name).join(MenuItem, MenuItem.id == OrderItem.menu_item_id)
                              .where(OrderItem.ended_at >= start, OrderItem.ended_at < end)).all())
    shifts_by_id = {s.id: s for s in db.scalars(select(Shift))}
    by = {k: defaultdict(lambda: [Decimal(0), Decimal(0)]) for k in ("type", "ingredient", "shift", "person", "dish")}
    on_time = Decimal(0)
    late = Decimal(0)
    for t, ing in txns:
        v = -t.value
        q = -t.qty
        if t.late:
            late += v
        else:
            on_time += v
        wt = t.waste_type or "khac"
        by["type"][d.WASTE_TYPES.get(wt, ("", wt))[1] + (" (ghi muộn)" if t.late else "")][0] += v
        by["ingredient"][ing.name][0] += v
        by["ingredient"][ing.name][1] += q
        if t.shift_id and t.shift_id in shifts_by_id:
            s = shifts_by_id[t.shift_id]
            by["shift"][f"{s.business_date:%d/%m} {s.name}"][0] += v
        person = t.responsible_user_id if t.late else t.created_by
        if person:
            by["person"][names.get(person, "?")][0] += v
        slip = slips.get(t.source_id) if t.source_type == "waste_slip" else None
        if slip and slip.order_item_id and slip.order_item_id in dish_of:
            by["dish"][dish_of[slip.order_item_id]][0] += v

    def rows(m):
        return sorted(({"key": k, "value": money(v[0]), "qty": q3(v[1])} for k, v in m.items()), key=lambda r: -r["value"])

    return {
        "on_time_value": money(on_time), "late_value": money(late),
        "by_type": rows(by["type"]), "by_ingredient": rows(by["ingredient"]), "by_shift": rows(by["shift"]),
        "by_person": rows(by["person"]), "by_dish": rows(by["dish"]),
    }


def shift_report(db: Session, date_from: date, date_to: date) -> list[dict]:
    """Chênh lệch theo ca và theo bếp trưởng ca."""
    names = _names(db)
    out = []
    for s in db.scalars(select(Shift).where(Shift.business_date >= date_from, Shift.business_date <= date_to)
                        .order_by(Shift.planned_start.desc())):
        lines = db.scalars(select(CountLine).where(CountLine.shift_id == s.id, CountLine.location == d.BEP,
                                                   CountLine.variance_qty.is_not(None))).all()
        handed = dec(db.scalar(
            select(func.coalesce(func.sum(StockTxn.value), 0)).where(StockTxn.shift_id == s.id,
                                                                     StockTxn.txn_type == d.TXN_HANDOVER_IN)))
        invs = db.scalar(select(func.count()).select_from(Investigation).where(Investigation.shift_id == s.id)) or 0
        sold = db.scalar(select(func.coalesce(func.sum(OrderItem.qty), 0)).where(
            OrderItem.shift_id == s.id, OrderItem.status.in_([d.ITEM_COOKING, d.ITEM_DONE]))) or 0
        out.append({
            "id": s.id, "business_date": s.business_date, "name": s.name, "status": s.status,
            "opened_at": s.opened_at, "closed_at": s.closed_at, "planned_end": s.planned_end,
            "head_chef_id": s.head_chef_id, "head_chef": names.get(s.head_chef_id),
            "handed_value": money(handed), "variance_value": money(sum((ln.variance_value or 0 for ln in lines), Decimal(0))),
            "worst_tier": max((ln.tier for ln in lines), key=lambda t: [d.TIER_OK, d.TIER_WARN, d.TIER_MUST].index(t),
                              default=d.TIER_OK),
            "investigations": invs, "dishes": int(sold), "sod_violation": s.sod_violation, "sod_notes": s.sod_notes,
        })
    return out


def count_history(db: Session, date_from: date, date_to: date) -> list[dict]:
    """Lịch sử kiểm kê và cân trả: ai đếm, khi nào, lệch bao nhiêu."""
    start, end = _range(date_from, date_to)
    names = _names(db)
    out = []
    for ln, s, ing in db.execute(
        select(CountLine, CountSession, Ingredient).join(CountSession, CountSession.id == CountLine.session_id)
        .join(Ingredient, Ingredient.id == CountLine.ingredient_id)
        .where(CountLine.cutoff_at >= start, CountLine.cutoff_at < end).order_by(CountLine.cutoff_at.desc())
    ):
        out.append({
            "session_id": s.id, "code": s.code, "kind": d.COUNT_KINDS.get(s.kind, s.kind), "location": ln.location,
            "cutoff_at": ln.cutoff_at, "counted_by": names.get(s.counted_by), "ingredient": ing.name,
            "unit": ing.base_unit, "book_qty": ln.book_qty, "counted_qty": ln.counted_qty,
            "variance_qty": ln.variance_qty, "variance_pct": ln.variance_pct, "variance_value": ln.variance_value,
            "tier": ln.tier,
        })
    return out


def receipt_history(db: Session, date_from: date, date_to: date) -> list[dict]:
    start, end = _range(date_from, date_to)
    out = []
    for ln, r, ing in db.execute(
        select(ReceiptLine, Receipt, Ingredient).join(Receipt, Receipt.id == ReceiptLine.receipt_id)
        .join(Ingredient, Ingredient.id == ReceiptLine.ingredient_id)
        .where(Receipt.received_at >= start, Receipt.received_at < end).order_by(Receipt.received_at.desc())
    ):
        out.append({
            "receipt_id": r.id, "code": r.code, "supplier": r.supplier, "received_at": r.received_at,
            "status": r.status, "ingredient": ing.name, "unit": ln.unit, "ordered_qty": ln.ordered_qty,
            "received_qty": ln.received_qty, "unit_price": ln.unit_price, "unit_cost_base": ln.unit_cost_base,
            "base_unit": ing.base_unit,
        })
    return out


def sales_report(db: Session, date_from: date, date_to: date) -> list[dict]:
    """Số món bán, hủy, làm lại theo ngày / ca (số đối chiếu)."""
    start, end = _range(date_from, date_to)
    agg: dict[tuple, dict] = {}
    shifts_by_id = {s.id: s for s in db.scalars(select(Shift))}
    for oi in db.scalars(select(OrderItem).where(OrderItem.created_at.is_not(None))
                         .where((OrderItem.started_at >= start) | (OrderItem.ended_at >= start))
                         .where((OrderItem.started_at < end) | (OrderItem.ended_at < end))):
        t = oi.started_at or oi.ended_at
        if t is None or not (start <= t < end):
            continue
        s = shifts_by_id.get(oi.shift_id)
        k = (clock.business_date(t), s.name if s else "Ngoài ca")
        r = agg.setdefault(k, {"date": k[0], "shift": k[1], "sold": 0, "cancel_after_start": 0, "remade": 0,
                               "cancel_before_start": 0, "out_of_stock": 0})
        if oi.status in (d.ITEM_COOKING, d.ITEM_DONE):
            r["sold"] += oi.qty
        elif oi.status == d.ITEM_REMADE:
            r["remade"] += oi.qty
        elif oi.status == d.ITEM_CANCELLED:
            r["cancel_after_start" if oi.started_at else "cancel_before_start"] += oi.qty
        elif oi.status == d.ITEM_OUT:
            r["out_of_stock"] += oi.qty
    return sorted(agg.values(), key=lambda r: (r["date"], r["shift"]), reverse=True)


def audit_list(db: Session, entity: str | None = None, actor_id: int | None = None, action: str | None = None,
               limit: int = 200) -> list[dict]:
    names = _names(db)
    stmt = select(AuditLog).order_by(AuditLog.id.desc()).limit(limit)
    if entity:
        stmt = stmt.where(AuditLog.entity == entity)
    if actor_id:
        stmt = stmt.where(AuditLog.actor_id == actor_id)
    if action:
        stmt = stmt.where(AuditLog.action.ilike(f"%{action}%"))
    return [{"id": a.id, "at": a.at, "actor": names.get(a.actor_id, "Hệ thống"), "action": a.action,
             "entity": a.entity, "entity_id": a.entity_id, "before": a.before, "after": a.after, "note": a.note}
            for a in db.scalars(stmt)]


def txn_list(db: Session, ingredient_id: int | None = None, location: str | None = None, txn_type: str | None = None,
             shift_id: int | None = None, limit: int = 300) -> list[dict]:
    """Sổ giao dịch kho: ai tạo, ai duyệt từng giao dịch (AC-18)."""
    names = _names(db)
    ings = {i.id: i for i in db.scalars(select(Ingredient))}
    stmt = select(StockTxn).order_by(StockTxn.id.desc()).limit(limit)
    if ingredient_id:
        stmt = stmt.where(StockTxn.ingredient_id == ingredient_id)
    if location:
        stmt = stmt.where(StockTxn.location == location)
    if txn_type:
        stmt = stmt.where(StockTxn.txn_type == txn_type)
    if shift_id:
        stmt = stmt.where(StockTxn.shift_id == shift_id)
    return [{
        "id": t.id, "occurred_at": t.occurred_at, "type": t.txn_type, "type_label": d.TXN_TYPES.get(t.txn_type, t.txn_type),
        "ingredient": ings[t.ingredient_id].name, "unit": ings[t.ingredient_id].base_unit, "location": t.location,
        "qty": t.qty, "value": t.value, "shift_id": t.shift_id, "source_type": t.source_type, "source_id": t.source_id,
        "ref_txn_id": t.ref_txn_id, "waste_type": t.waste_type, "late": t.late, "cause": t.cause,
        "responsible": names.get(t.responsible_user_id), "created_by": names.get(t.created_by),
        "approved_by": names.get(t.approved_by), "note": t.note,
    } for t in db.scalars(stmt)]


def handover_qty(db: Session, shift_id: int) -> Decimal:
    return dec(db.scalar(select(func.coalesce(func.sum(HandoverLine.issued_qty), 0))
                         .join(Handover, Handover.id == HandoverLine.handover_id)
                         .where(Handover.shift_id == shift_id, Handover.status == "signed")))
