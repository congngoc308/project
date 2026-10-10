"""AC-12 → AC-18: phản ứng khi hết hàng, tất toán, xu hướng, báo cáo, phân quyền, nhật ký."""

import io
from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy import select

from app import clock
from app.models import Alert, StockTxn
from tests.helpers import advance, current_shift, handover, receive, return_shift, sell


def _open_shift_with_beef(as_user, ref, beef=1000, rau=3000):
    sk, chef = as_user("thukho1"), as_user("beptruong1")
    receive(sk, ref["BO"], beef / 1000)
    receive(sk, ref["RAU"], rau / 1000, price=30000)
    advance()
    handover(sk, chef, ref["beptruong1"], "2221", {ref["BO"]: beef, ref["RAU"]: rau})
    advance()


def test_ac12_ingredient_out_locks_dishes_moves_waiting_items_alerts(as_user, ref, db):
    _open_shift_with_beef(as_user, ref, beef=1000)
    sv, k, mgr = as_user("phucvu1"), as_user("bep1"), as_user("quanly")
    sell(sv, k, ref["BX_dish"], 5)
    waiting = sv.post("/orders", {"table_label": "B7", "items": [
        {"menu_item_id": ref["BX_dish"]}, {"menu_item_id": ref["GR_dish"]}]})
    advance()

    res = k.post("/kitchen/out-of-ingredient", {"ingredient_id": ref["BO"]})
    assert "variance_qty" not in res  # bếp không thấy tồn sổ / chênh lệch
    assert set(res["locked_menu_items"]) == {ref["BX_dish"], ref["PB_dish"], ref["CC_dish"]}
    menu = {m["id"]: m for m in sv.get("/menu")}
    assert not menu[ref["BX_dish"]]["sellable"] and menu[ref["GR_dish"]]["sellable"]
    sv.post("/orders", {"table_label": "B8", "items": [{"menu_item_id": ref["PB_dish"]}]}, expect=409)

    notes = sv.get("/server/notifications")
    assert [(n["table"], n["name"]) for n in notes] == [("B7", "Bò xào")]
    statuses = {i["name"]: i["status"] for i in sv.get(f"/orders/{waiting['id']}")["items"]}
    assert statuses == {"Bò xào": "out_of_stock", "Gà rang": "waiting"}

    alert = db.scalar(select(Alert).where(Alert.kind == "out_with_book"))
    assert alert and alert.severity == "critical"
    inv = mgr.get("/investigations?status=open")[0]
    assert float(inv["variance_qty"]) == 900.0

    # Phiếu mở từ báo hết: phải kiểm kê xác nhận trước khi kết luận.
    mgr.post(f"/investigations/{inv['id']}/conclude", {"note": "x", "allocations": [
        {"cause": "that_thoat_ca", "qty": 900}]}, expect=409)
    # Bán lại cần bằng chứng có hàng.
    as_user("beptruong1").post(f"/menu-items/{ref['BX_dish']}/resell", {"note": ""}, expect=409)
    detail = mgr.post(f"/investigations/{inv['id']}/verify", {"found_qty": 300, "note": "Còn 300 g trong tủ mát"})
    assert float(detail["settled_qty"]) == 300.0 and detail["status"] == "open"
    mgr.post(f"/investigations/{inv['id']}/conclude", {"note": "Không giải trình được phần còn lại", "allocations": [
        {"cause": "khong_xac_dinh", "qty": 600}]})
    as_user("beptruong1").post(f"/menu-items/{ref['BX_dish']}/resell", {"note": "Tìm thấy hàng"})
    assert {m["id"]: m for m in sv.get("/menu")}[ref["BX_dish"]]["sellable"]


def test_ac13_receipt_error_conclusion_is_linked_adjustment_not_waste(as_user, ref, db):
    sk, chef, mgr = as_user("thukho1"), as_user("beptruong1"), as_user("quanly")
    r = receive(sk, ref["BO"], 1)  # phiếu ghi 1 kg nhưng hóa đơn chỉ 1 lạng
    advance()
    out = handover(sk, chef, ref["beptruong1"], "2221", {ref["BO"]: 1000}, received={ref["BO"]: 100}, expect=409)
    assert out["response"]["data"][0]["received_qty"] == 100
    sk.post(f"/handovers/{out['handover']['id']}/cancel", {"note": "Cân lại thấy chỉ có 100 g"})
    handover(sk, chef, ref["beptruong1"], "2221", {ref["BO"]: 100})
    advance(60)
    s = sk.post("/counts", {"kind": "periodic", "location": "kho", "ingredient_ids": [ref["BO"]]})
    sk.post(f"/counts/{s['id']}/submit", {"counts": {str(ref["BO"]): 0}})

    inv = mgr.get(f"/investigations/{mgr.get('/investigations?status=open')[0]['id']}")
    assert inv["location"] == "kho" and float(inv["variance_qty"]) == 900.0
    assert inv["suggestions"][0]["cause"] == "nhap_sai"
    waste_before = mgr.get("/reports/waste")
    mgr.post(f"/investigations/{inv['id']}/conclude", {"note": "Hóa đơn 1 lạng, nhập nhầm 1 kg", "allocations": [
        {"cause": "nhap_sai", "qty": 900, "receipt_line_id": r["lines"][0]["id"]}]})

    waste_after = mgr.get("/reports/waste")
    assert waste_after["late_value"] == waste_before["late_value"]
    assert waste_after["on_time_value"] == waste_before["on_time_value"]
    linked = db.scalar(select(StockTxn).where(StockTxn.txn_type == "linked_adjustment"))
    assert float(linked.qty) == -900.0 and linked.ref_txn_id == db.get(StockTxn, linked.ref_txn_id).id
    assert db.get(StockTxn, linked.ref_txn_id).txn_type == "receipt"
    assert linked.responsible_user_id == ref["thukho1"]
    receipt = sk.get(f"/receipts/{r['id']}")
    assert float(receipt["lines"][0]["linked_adjustment"]) == -900.0
    avt = mgr.get("/reports/avt?location=kho")
    beef = next(x for x in avt["rows"] if x["ingredient"] == "Thịt bò")
    assert beef["explained"][0]["cause"] == "nhap_sai"


def test_ac14_dish_out_only_locks_sale(as_user, ref, db):
    k, sv = as_user("bep1"), as_user("phucvu1")
    before = db.scalars(select(StockTxn)).all()
    k.post("/menu/out-of-dish", {"menu_item_id": ref["CC_dish"], "reason": "Hỏng chảo chiên"})
    assert not {m["id"]: m for m in sv.get("/menu")}[ref["CC_dish"]]["sellable"]
    assert db.scalars(select(StockTxn)).all() == before
    sv.post("/menu-items/%d/resell" % ref["CC_dish"], {"note": ""}, expect=403)
    as_user("beptruong1").post(f"/menu-items/{ref['CC_dish']}/resell", {"note": "Đã có chảo"})


def test_ac15_small_same_direction_shortage_three_shifts_trend(as_user, ref, db):
    sk, sv, k = as_user("thukho1"), as_user("phucvu1"), as_user("bep1")
    chefs = [("beptruong1", "2221", 9, 6), ("beptruong2", "2222", 9, 14), ("beptruong1", "2221", 10, 6)]
    for i, (chef, pin, day, hour) in enumerate(chefs):
        clock.set_fixed(datetime(2026, 10, day, hour, 30, tzinfo=clock.VN))
        receive(sk, ref["BO"], 1)
        receive(sk, ref["RAU"], 2, price=30000)
        advance()
        handover(sk, as_user(chef), ref[chef], pin, {ref["BO"]: 1000, ref["RAU"]: 2000})
        advance()
        sell(sv, k, ref["BX_dish"], 49, table=f"B{i}")
        advance(30)
        return_shift(sk, current_shift(sk)["id"], {ref["BO"]: 0, ref["RAU"]: 2000 - 49 * 30})
        assert db.scalar(select(Alert).where(Alert.kind.like("variance_%"))) is None
        found = db.scalar(select(Alert).where(Alert.kind == "trend"))
        assert (found is not None) == (i == 2)
    dash = as_user("quanly").get("/dashboard")
    assert any(a["kind"] == "trend" and "Thịt bò" in a["title"] for a in dash["alerts"])


def test_ac16_avt_report_units_money_sorting_filters_excel(as_user, ref):
    _open_shift_with_beef(as_user, ref, beef=1000, rau=1000)
    sk, mgr = as_user("thukho1"), as_user("quanly")
    sell(as_user("phucvu1"), as_user("bep1"), ref["BX_dish"], 5)
    advance(30)
    return_shift(sk, current_shift(sk)["id"], {ref["BO"]: 0, ref["RAU"]: 700})  # rau thiếu 150 g
    rep = mgr.get("/reports/avt?location=bep")
    names = [r["ingredient"] for r in rep["rows"]]
    assert names[:2] == ["Thịt bò", "Rau cải"]  # sắp theo tiền: 225.000 đ > 4.500 đ
    beef = rep["rows"][0]
    assert (float(beef["theoretical"]), float(beef["actual"]), float(beef["unexplained"])) == (100.0, 1000.0, 900.0)
    assert float(beef["pct"]) == 900.0 and float(beef["value"]) == 225000.0 and beef["tier"] == "must"
    only = mgr.get(f"/reports/avt?ingredient_id={ref['RAU']}&by_shift=true")
    assert [r["ingredient"] for r in only["rows"]] == ["Rau cải"] and only["rows"][0]["shift"]
    kho = mgr.get("/reports/avt?location=kho")["rows"]
    assert [r["ingredient"] for r in kho] == ["Dầu ăn"] and float(kho[0]["theoretical"]) == 25.0
    assert float(kho[0]["unexplained"]) == 0.0  # nhóm C chưa kiểm kê kỳ này

    xlsx = mgr._call("GET", "/reports/avt.xlsx?location=bep", 200)
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws["A4"].value == "Thịt bò" and ws["L4"].value == 225000


def test_ac17_permissions(as_user, ref, db):
    _open_shift_with_beef(as_user, ref)
    sk = as_user("thukho1")
    sell(as_user("phucvu1"), as_user("bep1"), ref["BX_dish"], 5)
    advance(30)
    return_shift(sk, current_shift(sk)["id"], {ref["BO"]: 0, ref["RAU"]: 3000 - 150})
    inv = as_user("quanly").get("/investigations?status=open")[0]

    # Bếp trưởng đã ký nhận ca không thể kết luận phiếu của ca đó, kể cả khi kiêm quản lý.
    body = {"note": "Tự kết luận", "allocations": [{"cause": "khong_xac_dinh", "qty": 900}]}
    as_user("beptruong1").post(f"/investigations/{inv['id']}/conclude", body, expect=403)
    mgr_chef = as_user("quanly")
    from app.models import Shift, User

    with db.begin():
        db.get(User, ref["quanly"]).roles = ["manager", "head_chef"]
        db.get(Shift, inv["shift"]["id"]).head_chef_id = ref["quanly"]
    mgr_chef.post(f"/investigations/{inv['id']}/conclude", body, expect=403)

    # Không đủ quyền thì không sửa công thức, không duyệt điều chỉnh / hao hụt của chính mình.
    as_user("beptruong1").post(f"/menu-items/{ref['BX_dish']}/recipes", {"reason": "x", "lines": [
        {"ingredient_id": ref["BO"], "qty": 10}]}, expect=403)
    as_user("thukho1").post(f"/investigations/{inv['id']}/conclude", body, expect=403)
    slip = as_user("bep1").post("/waste", {"ingredient_id": ref["BO"], "qty": 10, "waste_type": "do_vo_roi",
                                           "reason": "Rơi"})
    as_user("bep1").post(f"/waste/{slip['id']}/approve", expect=403)
    as_user("bep2").post(f"/waste/{slip['id']}/approve", expect=403)
    assert as_user("beptruong2").post(f"/waste/{slip['id']}/approve")["status"] == "approved"
    big = as_user("bep1").post("/waste", {"ingredient_id": ref["BO"], "qty": 100, "waste_type": "het_han",
                                          "reason": "Hết hạn"})
    assert big["status"] == "pending_manager"
    as_user("beptruong2").post(f"/waste/{big['id']}/approve", expect=403)


def test_ac18_who_created_changed_approved_cancelled(as_user, ref):
    _open_shift_with_beef(as_user, ref)
    sv, k = as_user("phucvu1"), as_user("bep1")
    o = sv.post("/orders", {"table_label": "B2", "items": [{"menu_item_id": ref["BX_dish"]}]})
    sv.patch(f"/order-items/{o['items'][0]['id']}", {"note": "Ít cay"})
    sv.post(f"/order-items/{o['items'][0]['id']}/cancel", {"reason": "Khách đổi món"})
    mgr = as_user("quanly")
    txns = mgr.get(f"/txns?ingredient_id={ref['BO']}&txn_type=handover_in")
    assert txns[0]["created_by"] == "Phạm Thị Thu" and txns[0]["approved_by"] == "Nguyễn Văn Bình"
    actions = {(a["action"], a["actor"]) for a in mgr.get("/audit")}
    assert ("order.create", "Đặng Thu Trang") in actions
    assert ("order_item.update", "Đặng Thu Trang") in actions
    assert ("order_item.cancel", "Đặng Thu Trang") in actions
    assert ("handover.sign", "Nguyễn Văn Bình") in actions
    assert ("receipt.receive", "Phạm Thị Thu") in actions
    k.get("/audit", expect=403)
